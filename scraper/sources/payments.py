"""Digital payments (RBI Payment System Indicators) and GST administrative data.

RBI publishes one "Payment System Indicators" page per month at
https://www.rbi.org.in/Scripts/PSIUserView.aspx?Id=N (Id 1 = June 2021, +1 per month).  Each page
carries three months for every row: the current month, the previous month and the same month a
year earlier.  A one-off backfill of every second page (Id 2, 4, ... = 32 requests) is stored as a
seed (seed/rbi-psi.json); production runs fetch only the latest few pages and merge.
"""

from __future__ import annotations

import re
import time
from datetime import datetime, timezone

from bs4 import BeautifulSoup

from .base import (
    clean_pts, iso, load_seed, merge_pts, month_ts, record, register, save_seed, series, ts, view,
    yoy_pts,
)

SECTION = "digital-administrative"
RBI_PSI_LIST = "https://www.rbi.org.in/Scripts/PSIUserView.aspx"
RBI_PSI_PAGE = "https://www.rbi.org.in/Scripts/PSIUserView.aspx?Id={id}"
PSI_FIRST_ID_MONTH = (2021, 6)  # Id 1

PROBE_URLS = [
    RBI_PSI_LIST,
    RBI_PSI_PAGE.format(id=64),
]

MONTHS = ["january", "february", "march", "april", "may", "june", "july", "august", "september",
          "october", "november", "december"]


def _sleep():
    time.sleep(1.0)


# ------------------------------------------------------------------ RBI PSI parsing

def _norm_label(label: str):
    """'2.6 UPI @' -> ('2.6', 'upi');  'Total - Retail Payments (2+3+4+5+6)' -> ('', 'total retail payments')."""
    s = re.sub(r"\s+", " ", label.replace("\xa0", " ")).strip()
    s = re.sub(r"\(\s*[\d+\s]+\)\s*$", "", s).strip()   # trailing "(2+3+4+5+6)" of the total rows
    m = re.match(r"^(\d+(?:\.\d+)*)\.?\s+(.*)$", s)
    code, name = (m.group(1), m.group(2)) if m else ("", s)
    name = re.sub(r"[\s$@*#^]+$", "", name)
    name = re.sub(r"[^a-z0-9 ]+", " ", name.lower())
    name = re.sub(r"\s+", " ", name).strip()
    return code, name


def _key(label: str) -> str:
    code, name = _norm_label(label)
    return f"{code} {name}".strip()


def _month_add(y: int, m: int, k: int):
    n = y * 12 + (m - 1) + k
    return n // 12, n % 12 + 1


def parse_psi_page(html: str):
    """Parse one RBI PSI page. Returns (year, month, {key: {"v": {ts: vol_lakh}, "x": {ts: value_crore}}}).

    Only the current and the previous month are returned (the year-ago column is ignored: later
    pages carry the more recent revision of the same month)."""
    soup = BeautifulSoup(html, "lxml")
    text = soup.get_text(" ", strip=True)
    m = re.search(r"Payment System Indicators\s*-\s*([A-Za-z]+)\s+(20\d\d)", text)
    if not m or m.group(1).lower() not in MONTHS:
        raise ValueError("PSI page: title with month not found")
    mon, yr = MONTHS.index(m.group(1).lower()) + 1, int(m.group(2))
    if "Volume (lakh)" not in text or "Value (" not in text or "crore" not in text:
        raise ValueError("PSI page: units header (Volume (lakh) / Value crore) not found")
    t_cur = month_ts(yr, mon)
    t_prev = month_ts(*_month_add(yr, mon, -1))

    out = {}
    header_seen = 0
    for tr in soup.find_all("tr"):
        cells = [re.sub(r"\s+", " ", c.get_text(" ", strip=True)) for c in tr.find_all(["td", "th"], recursive=False)]
        if len(cells) >= 8 and cells[0].startswith("FY") and len(cells) == 8:
            # header row of a data table: ['FY 2025-26', '2025 September', '2026', '2026', ...]
            hm = re.match(r"^(\d{4})\s+([A-Za-z]+)$", cells[1])
            if not hm or int(hm.group(1)) != yr - 1 or hm.group(2).lower() != MONTHS[mon - 1]:
                raise ValueError(f"PSI page: unexpected year-ago header {cells[1]!r} for {mon}/{yr}")
            header_seen += 1
            continue
        if len(cells) == 9 and cells[0]:
            from .base import _num
            nums = [_num(c) for c in cells[1:]]
            if sum(n is not None for n in nums) < 6:
                continue
            k = _key(cells[0])
            if k in out:
                continue  # first occurrence wins (e.g. repeated rows)
            nums = [0 if n is None else n for n in nums]
            out[k] = {"v": {t_prev: nums[2], t_cur: nums[3]}, "x": {t_prev: nums[6], t_cur: nums[7]}}
    if header_seen < 2:
        raise ValueError(f"PSI page {mon}/{yr}: expected 2 data-table headers, saw {header_seen}")
    if "2 credit transfers retail" not in out:
        raise ValueError(f"PSI page {mon}/{yr}: retail credit transfers row missing; layout changed")
    if "2.6 upi" not in out:
        raise ValueError(f"PSI page {mon}/{yr}: UPI row missing")
    return yr, mon, out


def _fetch_psi(http, page_id: int):
    html = http.text(RBI_PSI_PAGE.format(id=page_id))
    _sleep()
    return parse_psi_page(html)


def _fold(acc: dict, parsed_out: dict, newer_wins=True):
    for k, d in parsed_out.items():
        a = acc.setdefault(k, {"v": {}, "x": {}})
        for kind in ("v", "x"):
            a[kind].update(d[kind])


def backfill_psi(http):
    """One-off: fetch every second PSI page (Id 2,4,...), save to seed/rbi-psi.json."""
    html = http.text(RBI_PSI_LIST)
    ids = [int(i) for i in re.findall(r"PSIUserView\.aspx\?Id=(\d+)", html)]
    top = max(ids)
    acc: dict = {}
    for pid in range(2, top + 1, 2):
        _fold(acc, _fetch_psi(http, pid)[2])
    if top % 2 == 1:  # make sure the newest page is included
        _fold(acc, _fetch_psi(http, top)[2])
    seed = {k: {"v": sorted([t, v] for t, v in d["v"].items()), "x": sorted([t, v] for t, v in d["x"].items())}
            for k, d in acc.items()}
    save_seed("rbi-psi", seed)
    return seed


_PSI_CACHE: dict = {}


def psi_data(http) -> dict:
    """{key: {'v': pts, 'x': pts}} = seed merged with the latest pages from RBI (cached per process)."""
    if "d" in _PSI_CACHE:
        return _PSI_CACHE["d"]
    seed = load_seed("rbi-psi")
    if not seed:
        raise RuntimeError("seed/rbi-psi.json missing: run backfill_psi once")
    html = http.text(RBI_PSI_LIST)
    _sleep()
    ids = sorted({int(i) for i in re.findall(r"PSIUserView\.aspx\?Id=(\d+)", html)})
    if len(ids) < 6:
        raise RuntimeError("PSI listing: too few pages found; layout changed")
    live: dict = {}
    latest = None
    for pid in ids[-3:]:  # oldest first so the newest revision wins
        y, m, out = _fetch_psi(http, pid)
        _fold(live, out)
        latest = (y, m)
    # the expected month of the newest Id
    exp = _month_add(*PSI_FIRST_ID_MONTH, ids[-1] - 1)
    if latest != exp:
        raise RuntimeError(f"PSI newest page is {latest}, expected {exp} for Id {ids[-1]}")
    merged = {}
    for k in set(seed) | set(live):
        s = seed.get(k, {"v": [], "x": []})
        lv = live.get(k, {"v": {}, "x": {}})
        merged[k] = {
            "v": merge_pts(s["v"], sorted([t, v] for t, v in lv["v"].items())),
            "x": merge_pts(s["x"], sorted([t, v] for t, v in lv["x"].items())),
        }
    now = datetime.now(timezone.utc).timestamp()
    last_t = merged["2.6 upi"]["v"][-1][0]
    if now - last_t > 125 * 86400:
        raise RuntimeError(f"PSI data stale: latest month {iso(last_t)}")
    _PSI_CACHE["d"] = merged
    return merged


def _get(http, key: str, kind: str) -> list:
    d = psi_data(http)
    if key not in d:
        raise RuntimeError(f"PSI row {key!r} not found")
    pts = [p for p in d[key][kind]]
    if len(pts) < 24:
        raise RuntimeError(f"PSI row {key!r}/{kind}: only {len(pts)} points")
    return pts


def _check_series(name, pts, lo, hi, max_jump=3.0):
    vals = [v for _, v in pts]
    if not all(lo <= v <= hi for v in vals[-12:]):
        raise RuntimeError(f"{name}: recent values outside plausible range [{lo},{hi}]: {vals[-3:]}")
    for (t0, a), (t1, b) in zip(pts[-24:], pts[-23:]):
        if a > 0 and b > 0 and not (1 / max_jump < b / a < max_jump):
            raise RuntimeError(f"{name}: implausible jump {iso(t0)} {a} -> {iso(t1)} {b}")


def _payment_record(http, slug, title, description, rows, qna, vol_range, val_range, source_note="",
                    related=None, yoy_ids=None):
    """rows: [(series_id, series_name, psi_key)]; first row is the headline."""
    vol_s, val_s = [], []
    for sid, sname, key in rows:
        v = _get(http, key, "v")
        x = _get(http, key, "x")
        if sid == rows[0][0]:
            _check_series(sname + " volume", v, *vol_range)
            _check_series(sname + " value", x, *val_range)
        vol_s.append(series(sid + "_vol", sname, "Lakh transactions", v))
        val_s.append(series(sid + "_val", sname, "Rs. crore", x))
    yv = [series(s["id"] + "_yoy", s["name"], "Per cent", yoy_pts(s["pts"])) for s in vol_s]
    yx = [series(s["id"] + "_yoy", s["name"], "Per cent", yoy_pts(s["pts"])) for s in val_s]
    if yoy_ids is not None:
        yv = yv[:yoy_ids]
        yx = yx[:yoy_ids]
    return record(
        slug, title=title, description=description, frequency="M",
        views=[
            view("Volume", *vol_s, vid="volume"),
            view("Value", *val_s, vid="value"),
            view("Volume YoY %", *yv, vid="vol_yoy"),
            view("Value YoY %", *yx, vid="val_yoy"),
        ],
        source="Reserve Bank of India - Payment System Indicators",
        source_url=RBI_PSI_LIST, qna=qna, related=related or [],
        notes=source_note or None,
    )


PSI_NOTE = ("Source: RBI monthly Payment System Indicators (provisional, revised in later releases). Volume is in lakh "
            "(1 lakh = 100,000 = 0.1 million) transactions; value in Rs. crore (1 crore = 10 million). Published roughly "
            "a month after the reference month; history on the RBI site starts June 2021.")
COMMON_CAVEAT = ("RBI data are provisional and are revised in later releases; the dashboard keeps the latest revision "
                 "seen. Counts are transactions processed by the system, not unique users or goods purchased.")


# ------------------------------------------------------------------ indicators (RBI PSI)

@register("upi-transactions", section=SECTION, name="UPI Transactions")
def upi(http):
    return _payment_record(
        http, "upi-transactions", "UPI Transactions (RBI)",
        "Unified Payments Interface (UPI) transactions processed each month: volume in lakh transactions and value in Rs. crore, "
        "as reported in the RBI's Payment System Indicators. UPI is India's dominant retail digital payment rail (person-to-person "
        "and person-to-merchant), so its volume tracks small-ticket consumption and its value tracks larger transfers. It is "
        "published within about a month, well before GDP data, making it a useful real-time proxy for private consumption and "
        "the formalisation of retail spending. 1 lakh = 0.1 million.",
        [("upi", "UPI", "2.6 upi"), ("ussd", "of which USSD (*99#)", "2.6.1 of which ussd")],
        [
            {"question": "What is this series?",
             "answer": "It is the monthly total of transactions routed over UPI, the NPCI-run instant payment system, as reported to the RBI by banks and payment providers.^Volume is shown in lakh transactions (1 lakh = 0.1 million) and value in Rs. crore. The USSD line is the small part made through the feature-phone *99# channel."},
            {"question": "How should I read it?",
             "answer": "Volume growth reflects adoption and frequency of small payments; value growth additionally reflects the average ticket size. Year-on-year growth is the cleanest comparison because of day-count and festival effects: months differ in number of days (February is short) and the festive season lifts October-November.^Value per transaction (value divided by volume) falling over time means UPI is increasingly used for tiny payments."},
            {"question": "Caveats and revisions?",
             "answer": "RBI figures are provisional and revised in later monthly releases. UPI data here are RBI's compilation and can differ slightly from the figures NPCI publishes on its own site (which also counts some categories differently).^History on the RBI Payment System Indicators pages starts June 2021; earlier months are not on those pages. " + COMMON_CAVEAT},
        ],
        (1e5, 1e8), (1e5, 1e9), source_note=PSI_NOTE, yoy_ids=1)


@register("rtgs-payments", section=SECTION, name="RTGS Payments")
def rtgs(http):
    return _payment_record(
        http, "rtgs-payments", "RTGS Payments (RBI)",
        "Real Time Gross Settlement (RTGS) transactions each month: volume in lakh and value in Rs. crore, split into customer and "
        "interbank transactions. RTGS carries large-value payments (minimum Rs. 2 lakh) between businesses and banks, so its value "
        "is a high-frequency indicator of corporate and financial-market activity and liquidity conditions.",
        [("rtgs", "RTGS total", "1 credit transfers rtgs"), ("cust", "Customer transactions", "1.1 customer transactions"),
         ("ibk", "Interbank transactions", "1.2 interbank transactions")],
        [
            {"question": "What is this series?",
             "answer": "RTGS settles large-value payments one by one in real time in central bank money. The series gives monthly counts and values of all RTGS transactions and of customer-initiated versus interbank ones."},
            {"question": "How should I read it?",
             "answer": "Value is dominated by a small number of very large transfers and swings with quarter-end, month-end and fiscal-year-end cash management; compare year-on-year. Customer volume is steadier and rises with business activity."},
            {"question": "Caveats?",
             "answer": "Value in Rs. crore, volume in lakh transactions. Interbank RTGS value also reflects market operations and can be lumpy. " + COMMON_CAVEAT},
        ],
        (10, 1e5), (1e5, 1e9), source_note=PSI_NOTE, yoy_ids=1)


@register("neft-payments", section=SECTION, name="NEFT Payments")
def neft(http):
    return _payment_record(
        http, "neft-payments", "NEFT Payments (RBI)",
        "National Electronic Funds Transfer (NEFT) transactions each month: volume in lakh and value in Rs. crore. NEFT is the "
        "workhorse for mid-sized account-to-account transfers (salaries, supplier payments, bill settlement) and is a "
        "steady gauge of retail and small-business payment activity.",
        [("neft", "NEFT", "2.5 neft")],
        [
            {"question": "What is this series?",
             "answer": "Monthly count and value of NEFT credit transfers processed in the RBI's retail payment data. NEFT settles 24x7 in half-hourly batches."},
            {"question": "How should I read it?",
             "answer": "Use year-on-year growth. NEFT volume grows more slowly than UPI because small payments have moved to UPI, while value remains large because it carries bigger transfers."},
            {"question": "Caveats?",
             "answer": "Units are lakh transactions and Rs. crore. " + COMMON_CAVEAT},
        ],
        (100, 1e7), (1e5, 1e9), source_note=PSI_NOTE)


@register("imps-payments", section=SECTION, name="IMPS Payments")
def imps(http):
    return _payment_record(
        http, "imps-payments", "IMPS Payments (RBI)",
        "Immediate Payment Service (IMPS) transactions each month: volume in lakh and value in Rs. crore. IMPS is an instant, "
        "24x7 mobile/internet interbank transfer service that is increasingly overshadowed by UPI for small payments but still "
        "carries sizeable transfers.",
        [("imps", "IMPS", "2.3 imps")],
        [
            {"question": "What is this series?",
             "answer": "Monthly count and value of IMPS transfers, from the RBI Payment System Indicators."},
            {"question": "How should I read it?",
             "answer": "A slow trend, now mostly flat to modest growth, shows how payment volumes are migrating between rails; look at the share of IMPS in total digital payments alongside UPI rather than reading it as a standalone demand signal."},
            {"question": "Caveats?",
             "answer": "Units are lakh transactions and Rs. crore. " + COMMON_CAVEAT},
        ],
        (100, 1e7), (1e4, 1e8), source_note=PSI_NOTE)


@register("nach-payments", section=SECTION, name="NACH Payments")
def nach(http):
    return _payment_record(
        http, "nach-payments", "NACH Payments (RBI)",
        "National Automated Clearing House (NACH) bulk transactions: credit (salary, dividend, subsidy and refund payouts) and "
        "debit (EMI, SIP, insurance premium and utility mandates). NACH debit tracks loan instalments and recurring "
        "obligations; NACH credit tracks bulk payouts including government benefit transfers.",
        [("nachdr", "NACH Debit", "3.2 nach dr"), ("nachcr", "NACH Credit", "2.4 nach cr")],
        [
            {"question": "What is this series?",
             "answer": "Monthly volumes and values of NACH debits (mandate-based collections) and credits (bulk payouts) from the RBI. ECS was merged into NACH from January 2020."},
            {"question": "How should I read it?",
             "answer": "NACH Debit is a proxy for the stock of EMIs and recurring payments and rises with credit growth and mutual-fund SIP flows. Credit spikes around government transfers and dividend seasons, so year-on-year comparison is best."},
            {"question": "Caveats?",
             "answer": "Units are lakh transactions and Rs. crore. Bulk files are lumpy and sensitive to month-end and holiday timing. " + COMMON_CAVEAT},
        ],
        (100, 1e7), (1e4, 1e8), source_note=PSI_NOTE)


@register("card-payments", section=SECTION, name="Card Payments")
def cards(http):
    return _payment_record(
        http, "card-payments", "Card Payments - Credit & Debit (RBI)",
        "Card transactions at merchants (PoS) and online/other channels, split into credit and debit cards, each month in lakh "
        "transactions and Rs. crore. Credit-card spend is a timely read on discretionary and big-ticket consumption; debit-card "
        "payments on everyday spending. Cash withdrawals at ATMs are excluded (see the ATM indicator).",
        [("cards", "All card payments", "4 card payments"),
         ("cc", "Credit cards", "4.1 credit cards"), ("cc_pos", "Credit cards - PoS", "4.1.1 pos based"),
         ("cc_oth", "Credit cards - online/other", "4.1.2 others"),
         ("dc", "Debit cards", "4.2 debit cards"), ("dc_pos", "Debit cards - PoS", "4.2.1 pos based"),
         ("dc_oth", "Debit cards - online/other", "4.2.2 others")],
        [
            {"question": "What is this series?",
             "answer": "Monthly count and value of payment transactions made with credit and debit cards, split into physical point-of-sale and other (mainly online) usage."},
            {"question": "How should I read it?",
             "answer": "Credit-card value growth is a relatively timely indicator of urban consumption and of the festive season. A rising online share indicates e-commerce. Debit-card volumes have been declining as customers shift to UPI."},
            {"question": "Caveats?",
             "answer": "Excludes ATM withdrawals. The PoS/online split follows RBI classification and the 'others' group includes e-commerce and card-not-present usage. " + COMMON_CAVEAT},
        ],
        (1000, 1e6), (1e4, 1e8), source_note=PSI_NOTE, yoy_ids=3)


@register("netc-fastag", section=SECTION, name="NETC (FASTag) Transactions")
def netc(http):
    return _payment_record(
        http, "netc-fastag", "NETC / FASTag Transactions (RBI)",
        "National Electronic Toll Collection (NETC, i.e. FASTag) transactions linked to bank accounts: monthly volume in lakh and "
        "value in Rs. crore. Toll payments follow commercial and passenger road traffic, so this is a monthly indicator of road "
        "freight and mobility activity. Note the RBI series covers only NETC linked to bank accounts (not prepaid-wallet FASTag), "
        "so value is smaller than the total toll collection shown elsewhere on the dashboard.",
        [("netc", "NETC (linked to bank account)", "3.3 netc linked to bank account")],
        [
            {"question": "What is this series?",
             "answer": "Count and value of toll payments made through NETC debits to bank accounts, a part of the 'debit transfers and direct debits' group in RBI payment statistics."},
            {"question": "How should I read it?",
             "answer": "Volume is a clean proxy for vehicle movement on tolled highways; month-on-month changes reflect trading days, with seasonal strength in festive and harvest periods and weakness in monsoon months. Use year-on-year growth."},
            {"question": "Caveats?",
             "answer": "Value is reported in Rs. crore; fee rates change in April each year, so value growth includes tariff effects. Coverage is of transactions debited to bank accounts under NETC, which may differ from the NHAI-reported total toll collection. " + COMMON_CAVEAT},
        ],
        (10, 1e5), (10, 1e5), source_note=PSI_NOTE)


@register("prepaid-instruments", section=SECTION, name="Prepaid Payment Instruments (Wallets & Cards)")
def ppi(http):
    return _payment_record(
        http, "prepaid-instruments", "Prepaid Payment Instruments (RBI)",
        "Transactions made with prepaid payment instruments (mobile wallets and prepaid cards) each month, in lakh transactions "
        "and Rs. crore. Wallet usage reflects small-value payments in transport, food delivery and app ecosystems.",
        [("ppi", "All PPIs", "5 prepaid payment instruments"), ("wallets", "Wallets", "5.1 wallets"),
         ("ppicards", "PPI cards", "5.2 cards")],
        [
            {"question": "What is this series?",
             "answer": "Monthly volume and value of transactions using prepaid payment instruments, as reported by PPI issuers to the RBI."},
            {"question": "How should I read it?",
             "answer": "Growth shows how much small-ticket spending still flows through wallets and prepaid cards rather than directly through UPI; the value per transaction is small."},
            {"question": "Caveats?",
             "answer": "Wallet data were added to the RBI table from November 2019. " + COMMON_CAVEAT},
        ],
        (1000, 1e7), (1e4, 1e7), source_note=PSI_NOTE, yoy_ids=1)


@register("digital-payments-total", section=SECTION, name="Total Digital Payments")
def total_digital(http):
    return _payment_record(
        http, "digital-payments-total", "Total Digital Payments (RBI)",
        "Aggregate of all digital payment systems: RTGS, retail credit transfers (UPI, NEFT, IMPS, ...), debit transfers, cards and "
        "prepaid instruments, plus the broader total including paper instruments. Monthly volume in lakh and value in Rs. crore. A "
        "broad gauge of the transaction economy.",
        [("digital", "Total digital payments", "total digital payments"),
         ("retail", "Total retail payments", "total retail payments"),
         ("all", "Total payments", "total payments")],
        [
            {"question": "What is this series?",
             "answer": "RBI's headline aggregates: total digital payments (RTGS, credit/debit transfers, cards, PPIs), total retail payments (excluding RTGS, including paper-based) and total payments."},
            {"question": "How should I read it?",
             "answer": "Volume growth is dominated by UPI; value growth is dominated by RTGS and large transfers, so the two can diverge. Year-on-year growth smooths seasonality."},
            {"question": "Caveats?",
             "answer": "The set of included systems has widened over time (e.g. new inclusions from November 2019); the series is provisional. " + COMMON_CAVEAT},
        ],
        (1e4, 1e8), (1e6, 1e10), source_note=PSI_NOTE, yoy_ids=3)


@register("card-atm-mobile-channels", section=SECTION, name="ATM, Mobile & Internet Payment Channels")
def channels(http):
    return _payment_record(
        http, "card-atm-mobile-channels", "Payment Channels: ATM, Mobile and Internet Banking (RBI)",
        "Channel view of payments: cash withdrawals at ATMs, mobile-app payments and internet-banking payments each month, in lakh "
        "transactions and Rs. crore. ATM withdrawals are a (declining) cash-demand indicator; mobile app payments overlap with UPI/IMPS "
        "and show digital adoption.",
        [("atm", "Cash withdrawal at ATMs", "3 cash withdrawal at atms"),
         ("mobile", "Mobile payments (app based)", "1 mobile payments mobile app based"),
         ("internet", "Internet payments (netbanking)", "2 internet payments netbanking internet browser based")],
        [
            {"question": "What is this series?",
             "answer": "Monthly volumes and values by channel: ATM cash withdrawals, mobile-app-based payments and internet/netbanking payments. Channels overlap with the payment systems (a UPI payment is also a mobile payment), so they must not be added to the system totals."},
            {"question": "How should I read it?",
             "answer": "ATM withdrawal value growth is a rough guide to cash demand; its slow growth against digital growth shows the shift away from cash."},
            {"question": "Caveats?",
             "answer": "Channel rows overlap with the payment-system rows. " + COMMON_CAVEAT},
        ],
        (1000, 1e8), (1e5, 1e9), source_note=PSI_NOTE)


@register("ccil-settlement", section=SECTION, name="CCIL Settlement (G-Sec, Forex, Rupee Derivatives)")
def ccil(http):
    return _payment_record(
        http, "ccil-settlement", "CCIL-Cleared Settlement Volumes (RBI)",
        "Transactions cleared and settled through the Clearing Corporation of India (CCIL): government securities (outright, repo, "
        "tri-party repo), forex and rupee derivatives. Monthly volume in lakh and value in Rs. crore; a gauge of money-market, "
        "debt-market and foreign-exchange market activity.",
        [("ccil", "All CCIL operated systems", "1 ccil operated systems"),
         ("gsec", "Govt securities clearing", "1.1 govt securities clearing"),
         ("outright", "G-Sec outright", "1.1.1 outright"), ("repo", "G-Sec repo", "1.1.2 repo"),
         ("triparty", "Tri-party repo", "1.1.3 tri party repo"),
         ("forex", "Forex clearing", "1.2 forex clearing"),
         ("deriv", "Rupee derivatives", "1.3 rupee derivatives")],
        [
            {"question": "What is this series?",
             "answer": "The monthly count and value of trades settled through the CCIL's clearing and settlement systems, grouped by market segment."},
            {"question": "How should I read it?",
             "answer": "Repo and tri-party repo values reflect the liquidity management behaviour of banks and mutual funds; forex clearing tracks interbank currency trading activity (including RBI and customer-driven flows). They are market-activity measures, not consumption indicators."},
            {"question": "Caveats?",
             "answer": "Volumes are in lakh trades and may be small numbers (e.g. 0.5 lakh) for the debt segments; check units. " + COMMON_CAVEAT},
        ],
        (0.1, 1e4), (1e5, 1e10), source_note=PSI_NOTE, yoy_ids=1)


# ================================================================== GST statistics (gst.gov.in)
# https://www.gst.gov.in/download/gststatistics links to Excel files on tutorial.gst.gov.in:
#   ewb-data-<FY>.xlsx                 state-wise e-way bill statistics by month (FY 2018-19 onwards)
#   statewise_GST_collection_<FY>.xlsx all-India sheet: CGST/SGST/IGST collection, refunds, net by month
#   GSTR-3B-<FY>.xlsx                  state-wise GSTR-3B filing summary by return period
import io

import openpyxl

GST_STATS_PAGE = "https://www.gst.gov.in/download/gststatistics"
GST_FILE_BASE = "https://tutorial.gst.gov.in/offlineutilities/gst_statistics/"

PROBE_URLS += [
    GST_STATS_PAGE,
    GST_FILE_BASE + "ewb-data-2025-26.xlsx",
    GST_FILE_BASE + "statewise_GST_collection_2025-26.xlsx",
    GST_FILE_BASE + "GSTR-3B-2025-2026.xlsx",
]

_FILE_CACHE: dict = {}


def _gst_links(http) -> dict:
    """{kind: {start_year: url}} discovered on the GST statistics page."""
    if "links" in _FILE_CACHE:
        return _FILE_CACHE["links"]
    html = http.text(GST_STATS_PAGE)
    _sleep()
    hrefs = re.findall(r'href="([^"]*gst_statistics/[^"]+\.xlsx)"', html)
    out = {"ewb": {}, "coll": {}, "g3b": {}}
    for h in hrefs:
        url = "https:" + h if h.startswith("//") else h
        name = url.rsplit("/", 1)[1]
        m = re.match(r"ewb-data-(\d{4})-\d{2}\.xlsx$", name)
        if m:
            out["ewb"][int(m.group(1))] = url
        m = re.match(r"statewise_GST_collection_(\d{4})-\d{2}\.xlsx$", name)
        if m:
            out["coll"][int(m.group(1))] = url
        m = re.match(r"GSTR-3B-(\d{4})-\d{4}\.xlsx$", name)
        if m:
            out["g3b"][int(m.group(1))] = url
    for k, n in (("ewb", 6), ("coll", 8), ("g3b", 8)):
        if len(out[k]) < n:
            raise RuntimeError(f"GST statistics page: only {len(out[k])} {k} files found; layout changed")
    _FILE_CACHE["links"] = out
    return out


def _wb(http, url):
    if url not in _FILE_CACHE:
        b = http.bytes(url)
        _sleep()
        _FILE_CACHE[url] = openpyxl.load_workbook(io.BytesIO(b), data_only=True)
    return _FILE_CACHE[url]


def _fnum(v):
    from .base import _num
    return _num(v)


def _fresh(last_ts, days, what):
    if datetime.now(timezone.utc).timestamp() - last_ts > days * 86400:
        raise RuntimeError(f"{what} stale: latest {iso(last_ts)}")


# ---------------------------------------------------------------- e-way bills

def parse_ewb(wb):
    """-> {ts: dict(within_n, within_v, out_n, out_v, in_n, in_v, states)} summed over state rows."""
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    hdr = None
    for i, r in enumerate(rows[:10]):
        if r[1] and str(r[1]).strip().upper() == "STATE NAME" and str(r[2]).strip().lower() == "year" \
                and str(r[3]).strip().lower() == "month":
            hdr = i
            break
    if hdr is None:
        raise ValueError("e-way bill sheet: header row (STATE NAME / Year / Month) not found")
    h = [(str(c or "").upper().replace("\n", " ")) for c in rows[hdr]]
    h = [x.replace(" ", "") for x in h]
    if not (h[4].startswith(("WITHIN", "INTRA")) and h[7].startswith(("OUTGOING", "OUTWARD", "INTERSTATEOUT"))
            and h[10].startswith(("INCOMING", "INWARD", "INTERSTATEIN"))):
        raise ValueError(f"e-way bill sheet: unexpected column groups {h[4:11]}")
    sub = [(str(c or "").upper().replace(chr(10), "").replace(" ", "")) for c in rows[hdr + 1]]
    cnt = ("NO.OFEWB", "NO.OFEWAYBILLS")
    val = ("ASSESSABLE", "ASSETVALUE")
    if not (sub[5].startswith(cnt) and sub[6].startswith(val) and sub[8].startswith(cnt)
            and sub[9].startswith(val) and sub[11].startswith(cnt)):
        raise ValueError(f"e-way bill sheet: unexpected sub-headers {sub[4:13]}")
    acc: dict = {}
    for r in rows[hdr + 2:]:
        if r[1] is None or not isinstance(r[2], int) or not isinstance(r[3], int):
            continue
        t = month_ts(r[2], r[3])
        a = acc.setdefault(t, {"within_n": 0, "within_v": 0.0, "out_n": 0, "out_v": 0.0, "in_n": 0, "in_v": 0.0, "states": 0})
        vals = [_fnum(r[k]) or 0 for k in (5, 6, 8, 9, 11, 12)]
        for key, v in zip(("within_n", "within_v", "out_n", "out_v", "in_n", "in_v"), vals):
            a[key] += v
        a["states"] += 1
    if not acc:
        raise ValueError("e-way bill sheet: no data rows")
    for t, a in acc.items():
        if a["states"] < 30:
            raise ValueError(f"e-way bill {iso(t)}: only {a['states']} state rows")
    return acc


def ewb_monthly(http) -> dict:
    links = _gst_links(http)["ewb"]
    acc: dict = {}
    for fy in sorted(links):
        acc.update(parse_ewb(_wb(http, links[fy])))
    return acc


@register("gst-eway-bills", section=SECTION, name="GST E-Way Bills")
def gst_eway(http):
    acc = ewb_monthly(http)
    # October 2021 in the GSTN file is a source defect: Maharashtra, Gujarat and Tamil Nadu report ~0 within-state
    # bills (total 41 mn vs ~60-70 mn in the adjacent months). Dropped rather than shown as a collapse.
    m = acc.pop(month_ts(2021, 10), None)
    if m is not None and m["within_n"] + m["out_n"] > 55e6:
        raise RuntimeError("e-way bill Oct 2021 no longer looks defective; review the exclusion in gst_eway()")
    ks = sorted(acc)
    if not 0.6 < (acc[ks[-1]]["within_n"] + acc[ks[-1]]["out_n"]) / (acc[ks[-2]]["within_n"] + acc[ks[-2]]["out_n"]) < 1.6:
        raise RuntimeError("latest e-way bill month differs implausibly from the previous month (partial data?)")
    total_n = [[t, acc[t]["within_n"] + acc[t]["out_n"]] for t in ks]
    total_v = [[t, round(acc[t]["within_v"] + acc[t]["out_v"], 2)] for t in ks]
    within_n = [[t, acc[t]["within_n"]] for t in ks]
    inter_n = [[t, acc[t]["out_n"]] for t in ks]
    within_v = [[t, round(acc[t]["within_v"], 2)] for t in ks]
    inter_v = [[t, round(acc[t]["out_v"], 2)] for t in ks]
    if len(total_n) < 60 or ks[0] != month_ts(2018, 7):
        raise RuntimeError(f"e-way bill history too short / not starting Jul 2018: {len(total_n)} months from {iso(ks[0])}")
    _check_series("ewb count", total_n, 1e6, 1e9, max_jump=4.0)
    _check_series("ewb value", total_v, 1e5, 1e8, max_jump=4.0)
    _fresh(ks[-1], 110, "e-way bill data")
    return record(
        "gst-eway-bills", title="GST E-Way Bills (monthly)",
        description="E-way bills generated each month in India and their assessable value (Rs. crore), from GST Network statistics. An "
                    "e-way bill is required for most movement of goods worth over Rs. 50,000, so the count is a direct, "
                    "high-frequency read on goods movement and trade activity (domestic and inter-state), available about six weeks after "
                    "the month. Split into within-state and inter-state (outgoing) bills.",
        frequency="M",
        views=[
            view("E-way bills (count)", series("ewb_total", "E-way bills generated", "Number", total_n),
                 series("ewb_within", "Within-state", "Number", within_n),
                 series("ewb_inter", "Inter-state (outgoing)", "Number", inter_n), vid="count"),
            view("Assessable value", series("ewb_value", "Assessable value of e-way bills", "Rs. crore", total_v),
                 series("ewb_within_v", "Within-state", "Rs. crore", within_v),
                 series("ewb_inter_v", "Inter-state (outgoing)", "Rs. crore", inter_v), vid="value"),
            view("Count YoY %", series("ewb_total_yoy", "E-way bills generated", "Per cent", yoy_pts(total_n)),
                 series("ewb_within_yoy", "Within-state", "Per cent", yoy_pts(within_n)),
                 series("ewb_inter_yoy", "Inter-state (outgoing)", "Per cent", yoy_pts(inter_n)), vid="count_yoy"),
            view("Value YoY %", series("ewb_value_yoy", "Assessable value", "Per cent", yoy_pts(total_v)), vid="value_yoy"),
        ],
        source="GST Network (gst.gov.in GST Statistics - E-Way Bill)", source_url=GST_STATS_PAGE,
        qna=[
            {"question": "What is this series?",
             "answer": "The number of e-way bills generated in a month, summed over all States and UTs, and the assessable (taxable) value of the goods covered. Total = within-state bills + bills for goods going out to other states. 'Incoming from other states' is the same inter-state bills seen from the receiving state, so it is not added.^Counts are plain numbers (a recent month is above 100 million) and values are in Rs. crore. Monthly statistics start July 2018 (e-way bills became compulsory for inter-state movement in April 2018)."},
            {"question": "How should I read it?",
             "answer": "Year-on-year growth in bill counts tracks the physical movement of goods; value growth also reflects prices and the mix of goods. Seasonality is marked: March (year-end) is high and festival and monsoon months move the series, so compare the same month across years.^It is a coincident indicator for goods trade, manufacturing dispatches and wholesale activity and tends to move with GST collections and industrial output."},
            {"question": "Caveats and revisions?",
             "answer": "Thresholds, state rules for intra-state bills (phased in during 2018) and the mix of goods covered have changed over time, so growth in 2018-19 is partly a roll-out effect.^One consignment can generate more than one bill, so counts are not distinct shipments. October 2021 is omitted because the published file shows near-zero within-state bills for Maharashtra, Gujarat and Tamil Nadu that month (a reporting defect). Data are as published by GSTN and may be revised. The statistics files are monthly only; no daily series is published there."},
        ],
        notes="Monthly state-wise file summed over States/UTs; total = within-state + outgoing inter-state bills. These totals run about 0.5-1.4% below the headline figures GSTN announces for the same months, so the two compilations differ slightly. October 2021 is omitted because the source file is defective for that month.",
    )


# ---------------------------------------------------------------- GST collections (all-India, by tax head)

def _month_blocks(ws):
    """Locate month column blocks. Returns (rows, first_data_row_index, blocks) where
    blocks = [((y, m), {label: column_index})] using the date header row and the label row below it."""
    rows = list(ws.iter_rows(values_only=True))
    for i, r in enumerate(rows[:15]):
        dcols = [c for c, v in enumerate(r) if isinstance(v, datetime)]
        if len(dcols) >= 1 and i + 1 < len(rows):
            labs = rows[i + 1]
            blocks = []
            for k, c in enumerate(dcols):
                end = dcols[k + 1] if k + 1 < len(dcols) else len(labs)
                m = {}
                for cc in range(c, end):
                    lab = str(labs[cc] or "").strip().upper()
                    if lab in ("CGST", "SGST", "IGST", "CESS", "TOTAL"):
                        m[lab] = cc
                blocks.append(((r[c].year, r[c].month), m))
            return rows, i + 2, blocks
    raise ValueError("GST collection sheet: month header row not found")


def _tot_ok(vals: dict, what: str):
    """TOTAL must equal the sum of the tax heads (within rounding)."""
    if "TOTAL" in vals and abs(sum(v for k, v in vals.items() if k != "TOTAL") - vals["TOTAL"]) > max(5, 0.002 * abs(vals["TOTAL"])):
        raise ValueError(f"GST {what}: components do not add to TOTAL ({vals})")


def parse_gst_allindia(wb):
    """All-India sheet -> {(section_key): {head: {ts: value}}} for sections:
    dom, imp, gross, ref_dom, ref_exp, ref_tot, net. Heads CGST/SGST/IGST/CESS/TOTAL (CESS only if the sheet has it)."""
    name = next((n for n in wb.sheetnames if "all" in n.lower() and "india" in n.lower()), None)
    if name is None:
        raise ValueError("GST collection workbook: no All-India sheet")
    rows, start, blocks = _month_blocks(wb[name])
    wanted = {"A": {"domestic": "dom", "imports": "imp", "gross revenue": "gross"},
              "B": {"domestic": "ref_dom", "export refund": "ref_exp", "total refund": "ref_tot"}}
    out: dict = {}
    section = None
    for r in rows[start:]:
        lab = str(r[0] or "").strip()
        if not lab:
            continue
        low = lab.lower()
        if low.startswith("a. collection"):
            section = "A"
            continue
        if low.startswith("b. refund"):
            section = "B"
            continue
        key = None
        if low.startswith("net gst revenue"):
            key = "net"
        elif section in wanted:
            for pref, k in wanted[section].items():
                if low.startswith(pref):
                    key = k
                    break
        if key is None:
            continue
        for (y, m), cols in blocks:
            if "TOTAL" not in cols:
                continue
            vals = {h: _fnum(r[c]) for h, c in cols.items()}
            if vals.get("TOTAL") is None:
                continue
            vals = {h: (0 if v is None else v) for h, v in vals.items()}  # blank = nil (e.g. imports CGST/SGST)
            if key == "gross" and vals["TOTAL"] == 0:
                continue
            _tot_ok(vals, f"{key} {y}-{m}")
            for h, v in vals.items():
                out.setdefault(key, {}).setdefault(h, {})[month_ts(y, m)] = v
    for k in ("dom", "imp", "gross", "ref_dom", "ref_exp", "ref_tot", "net"):
        if k not in out:
            raise ValueError(f"GST all-India sheet: section {k!r} not found; layout changed")
    # drop months where gross total is zero / absent (future months of the current FY)
    live = {t for t, v in out["gross"]["TOTAL"].items() if v > 0}
    for k in out:
        for h in out[k]:
            out[k][h] = {t: v for t, v in out[k][h].items() if t in live}
    return out


def parse_gst_statesum(wb):
    """Statewise sheet summed over States/UTs -> {head: {ts: value}} (domestic gross collection)."""
    name = next((n for n in wb.sheetnames if n.lower().startswith("collections")), None)
    if name is None:
        raise ValueError("GST collection workbook: no Collections-Statewise sheet")
    rows, start, blocks = _month_blocks(wb[name])
    out: dict = {}
    nstates: dict = {}
    for r in rows[start:]:
        code = r[0]
        if not (isinstance(code, int) or (isinstance(code, str) and code.strip().isdigit())):
            continue
        for (y, m), cols in blocks:
            if "TOTAL" not in cols:
                continue
            t = month_ts(y, m)
            nstates[t] = nstates.get(t, 0) + 1
            for h, c in cols.items():
                out.setdefault(h, {}).setdefault(t, 0)
                out[h][t] += _fnum(r[c]) or 0
    if "TOTAL" not in out:
        raise ValueError("GST statewise sheet: no TOTAL column")
    for t, n in nstates.items():
        if n < 35:
            raise ValueError(f"GST statewise sheet: only {n} state rows")
    live = {t for t, v in out["TOTAL"].items() if v > 0}
    return {h: {t: round(v, 2) for t, v in d.items() if t in live} for h, d in out.items()}


def gst_collections(http) -> dict:
    """Merge all FY workbooks. Returns {key: {head: pts}} with 'dom' extended back to Jul 2017 by state sums."""
    if "coll" in _FILE_CACHE:
        return _FILE_CACHE["coll"]
    links = _gst_links(http)["coll"]
    allindia: dict = {}
    states: dict = {}
    for fy in sorted(links):
        wb = _wb(http, links[fy])
        if any("all" in n.lower() and "india" in n.lower() for n in wb.sheetnames):
            for key, heads in parse_gst_allindia(wb).items():
                for h, d in heads.items():
                    allindia.setdefault(key, {}).setdefault(h, {}).update(d)
        s = parse_gst_statesum(wb)
        for h, d in s.items():
            states.setdefault(h, {}).update(d)
    if not allindia:
        raise RuntimeError("no all-India GST sheets found")
    res = {k: {h: [[t, v] for t, v in sorted(d.items())] for h, d in heads.items()} for k, heads in allindia.items()}
    first_ai = min(t for t in allindia["dom"]["TOTAL"])
    # extend domestic gross collection back with state sums where no all-India sheet exists
    dom_ext = {}
    for h, d in states.items():
        dom_ext[h] = [[t, v] for t, v in sorted(d.items()) if t < first_ai] + res["dom"].get(h, [])
    res["dom_ext"] = dom_ext
    _FILE_CACHE["coll"] = res
    return res


@register("gst-revenue-components", section=SECTION, name="GST Revenue by Component")
def gst_components(http):
    c = gst_collections(http)
    gross, dom, imp = c["gross"], c["dom"], c["imp"]
    tot = gross["TOTAL"]
    last = tot[-1][0]
    _fresh(last, 100, "GST collection data")
    if len(tot) < 40:
        raise RuntimeError("GST all-India history too short")
    _check_series("gross GST", tot, 5e4, 5e5, max_jump=1.8)
    # cross-check: the gross total equals domestic + imports for each month
    d_tot, i_tot = dict(map(tuple, dom["TOTAL"])), dict(map(tuple, imp["TOTAL"]))
    for t, v in tot[-24:]:
        if abs(d_tot[t] + i_tot[t] - v) > max(5, 0.002 * v):
            raise RuntimeError(f"GST gross != domestic + imports for {iso(t)}")
    # IGST on imports as a percentage... (not shown). Net series
    def s(sid, name, pts):
        return series(sid, name, "Rs. crore", pts)

    gross_heads = [s("g_total", "Gross GST revenue", tot), s("g_cgst", "CGST", gross["CGST"]),
                   s("g_sgst", "SGST", gross["SGST"]), s("g_igst", "IGST (domestic + imports)", gross["IGST"])]
    if "CESS" in gross:
        gross_heads.append(s("g_cess", "Compensation cess (until it was discontinued)", gross["CESS"]))
    dom_ext = c["dom_ext"]
    dom_heads = [s("d_total", "Domestic gross revenue", dom_ext["TOTAL"]), s("d_cgst", "CGST", dom_ext["CGST"]),
                 s("d_sgst", "SGST", dom_ext["SGST"]), s("d_igst", "IGST (domestic)", dom_ext["IGST"])]
    if "CESS" in dom_ext:
        dom_heads.append(s("d_cess", "Cess (domestic)", dom_ext["CESS"]))
    views = [
        view("Gross revenue by tax head", *gross_heads, vid="gross"),
        view("Domestic gross revenue (longer history)", *dom_heads, vid="domestic"),
        view("Domestic vs imports", s("dom", "Domestic gross revenue", dom["TOTAL"]),
             s("imp", "Gross revenue on imports (IGST + cess collected by Customs)", imp["TOTAL"]),
             s("tot", "Gross GST revenue", tot), vid="dom_imp"),
        view("Refunds", s("ref_total", "Total refunds", c["ref_tot"]["TOTAL"]),
             s("ref_dom", "Domestic refunds (RFD-05)", c["ref_dom"]["TOTAL"]),
             s("ref_exp", "Export refunds (via ICEGATE)", c["ref_exp"]["TOTAL"]), vid="refunds"),
        view("Net revenue", s("net_total", "Net GST revenue", c["net"]["TOTAL"]), s("net_cgst", "CGST", c["net"]["CGST"]),
             s("net_sgst", "SGST", c["net"]["SGST"]), s("net_igst", "IGST", c["net"]["IGST"]), vid="net"),
        view("YoY %", series("y_gross", "Gross GST revenue", "Per cent", yoy_pts(tot)),
             series("y_dom", "Domestic gross", "Per cent", yoy_pts(dom["TOTAL"])),
             series("y_imp", "Imports gross", "Per cent", yoy_pts(imp["TOTAL"])),
             series("y_ref", "Total refunds", "Per cent", yoy_pts(c["ref_tot"]["TOTAL"])),
             series("y_net", "Net GST revenue", "Per cent", yoy_pts(c["net"]["TOTAL"])), vid="yoy"),
    ]
    return record(
        "gst-revenue-components", title="GST Revenue by Component: CGST, SGST, IGST, Domestic vs Imports, Refunds",
        description="Monthly GST collections broken down by tax head (CGST, SGST, IGST and, while it lasted, compensation cess), by domestic "
                    "versus imports, with refunds (domestic and export) and net revenue, in Rs. crore, from the GST Network all-India "
                    "statistics. The dashboard already carries total GST collections; the split shows whether growth comes from domestic "
                    "activity (domestic CGST/SGST) or from imports (IGST collected by Customs), and refunds show how much of gross "
                    "revenue is returned, notably to exporters.",
        frequency="M", views=views,
        source="GST Network (gst.gov.in GST Statistics - Tax Collection)", source_url=GST_STATS_PAGE,
        qna=[
            {"question": "What is this series?",
             "answer": "All-India monthly GST revenue by component as compiled by GSTN: gross collection (CGST, SGST, IGST), split into domestic (paid on the GST portal) and imports (IGST paid at Customs), the refunds disbursed (domestic RFD-05 refunds and export refunds through ICEGATE) and net revenue (gross minus refunds). Values are in Rs. crore.^The all-India sheets cover April 2021 onward; the 'domestic gross revenue (longer history)' view extends domestic collection back to July 2017 by summing the state-wise collection tables."},
            {"question": "How should I read it?",
             "answer": "Domestic CGST and SGST follow domestic consumption and business turnover with a lag (collections in a month relate to supplies in the previous month). Import IGST follows import values and the exchange rate, so it can diverge sharply from domestic revenue. Refund swings (particularly export refunds) change net revenue even when gross is steady. Use year-on-year growth to remove seasonality (March and the festive months are strong)."},
            {"question": "Caveats and revisions?",
             "answer": "Numbers are provisional and are revised slightly on finalisation. IGST in the all-India split is collected before it is apportioned (settled) between the Centre and States, so CGST and SGST are not the full Centre and State shares. Compensation cess is shown only for the periods where GSTN reports it; it is not in the later sheets after the cess was withdrawn, so the gross total here may differ from headline totals that include cess in earlier years and from the headline GST total series on the dashboard by small amounts."},
        ],
        notes="From GSTN all-India GST collection sheets (April 2021 onward) and state-wise tables (July 2017 onward for domestic).",
    )


# ---------------------------------------------------------------- GSTR-3B filings

def parse_g3b(wb):
    """-> {ts: {eligible, by_due, after_due, total}} from the 'Total' row of a GSTR-3B filing summary."""
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    hdr = None
    for i, r in enumerate(rows[:15]):
        if str(r[0] or "").strip().upper().startswith("STATE CD") and any(isinstance(v, datetime) for v in r):
            hdr = i
            break
    if hdr is None:
        raise ValueError("GSTR-3B sheet: header row not found")
    dcols = [(c, v) for c, v in enumerate(rows[hdr]) if isinstance(v, datetime)]
    sub = rows[hdr + 1]
    out: dict = {}
    total_row = next((r for r in rows[hdr + 2:] if str(r[0] or "").strip().lower() in ("total", "grand total")), None)
    if total_row is None:
        raise ValueError("GSTR-3B sheet: Total row not found")
    for c, d in dcols:
        labs = [str(sub[c + k] or "").lower() for k in range(4)]
        if not (labs[0].startswith("no. of tax payers eligible") and "due date" in labs[1] and "after" in labs[2]
                and labs[3].startswith("total returns")):
            raise ValueError(f"GSTR-3B sheet: unexpected sub-headers {labs}")
        vals = [_fnum(total_row[c + k]) for k in range(4)]
        if any(v is None for v in vals):
            continue
        if vals[3] <= 0:
            continue
        out[month_ts(d.year, d.month)] = {"eligible": vals[0], "by_due": vals[1], "after_due": vals[2], "total": vals[3]}
    if not out:
        raise ValueError("GSTR-3B sheet: no months")
    return out


@register("gstr3b-returns-filed", section=SECTION, name="GSTR-3B Returns Filed")
def gstr3b(http):
    links = _gst_links(http)["g3b"]
    acc: dict = {}
    for fy in sorted(links):
        acc.update(parse_g3b(_wb(http, links[fy])))
    ks = sorted(acc)
    if len(ks) < 60:
        raise RuntimeError(f"GSTR-3B history too short: {len(ks)} months")
    total = [[t, acc[t]["total"]] for t in ks]
    by_due = [[t, acc[t]["by_due"]] for t in ks]
    after = [[t, acc[t]["after_due"]] for t in ks]
    elig = [[t, acc[t]["eligible"]] for t in ks]
    _check_series("gstr3b total", total, 1e6, 3e7, max_jump=3.0)
    _fresh(ks[-1], 170, "GSTR-3B data")
    pct = [[t, round(100 * acc[t]["total"] / acc[t]["eligible"], 2)] for t in ks if acc[t]["eligible"]]
    return record(
        "gstr3b-returns-filed", title="GSTR-3B Returns Filed (by return period)",
        description="Number of GSTR-3B (monthly summary) returns filed for each return period, with the number of taxpayers eligible to file, "
                    "filings made by the due date and those filed late, from the GST Network. A gauge of compliance and of the size of the "
                    "active GST taxpayer base; changes in on-time filing signal stress in small business cash flows or administrative "
                    "changes. The count of eligible filers rises over time as new taxpayers register.",
        frequency="M",
        views=[
            view("Returns filed", series("g3b_total", "Total returns filed", "Number", total),
                 series("g3b_due", "Filed by due date", "Number", by_due),
                 series("g3b_after", "Filed after due date", "Number", after),
                 series("g3b_elig", "Taxpayers eligible to file", "Number", elig), vid="filed"),
            view("Filing rate", series("g3b_pct", "Returns filed as % of eligible", "Per cent", pct), vid="pct"),
            view("YoY %", series("g3b_yoy", "Total returns filed", "Per cent", yoy_pts(total)), vid="yoy"),
        ],
        source="GST Network (gst.gov.in GST Statistics - Returns)", source_url=GST_STATS_PAGE,
        qna=[
            {"question": "What is this series?",
             "answer": "For each return period (the month of supply), the all-India total of GSTR-3B returns filed and the number of taxpayers who were eligible to file. The Total row of GSTN's state-wise summary is used. Dates are the return period, not the filing month: returns for a month are filed mostly in the following month (due around the 20th).^Counts are plain numbers of returns (about 8-10 million a month recently)."},
            {"question": "How should I read it?",
             "answer": "The total filed grows slowly with the registered base. The share filed by the due date measures compliance discipline. Because the count 'filed after due date' keeps growing as late filers catch up, the most recent periods look worse than they will end up, so compare recent months only with similar vintages."},
            {"question": "Caveats and revisions?",
             "answer": "The file is a snapshot 'as on' the data-considered date, so recent months are incomplete and earlier months already include late filings. Quarterly (QRMP) filers file GSTR-3B monthly or quarterly depending on scheme, which lifts the number of eligible taxpayers and filings in some months (notably the end of each quarter); a shift in scheme rules changes counts without a change in activity. The gross revenue the returns generate is not in this file."},
        ],
        notes="Snapshot of filing counts as on the file date; recent return periods are not yet complete.",
    )
